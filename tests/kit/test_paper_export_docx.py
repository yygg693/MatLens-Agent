# -*- coding: utf-8 -*-
"""论文工坊 docx 导出器测试。

只用 stdlib + pytest + python-docx；不 import paper.export 门面（它依赖尚未完成的
export_text），直接 import matlens_kit.paper.export_docx / md_model / refs。无网络。
"""
from __future__ import annotations

import io
import zipfile

from docx import Document

from matlens_kit.paper.export_docx import export_docx
from matlens_kit.paper.md_model import parse_paper
from matlens_kit.paper.refs import ref_lines


# ---------------------------------------------------------------------------
# 合成论文：write 含 # 标题、## Abstract、**Keywords:**、## 1. Introduction、
# ### 1.1、pipe 表格、化学式；search refs 2 条。
# ---------------------------------------------------------------------------
WRITE_CONTENT = """# Ionic Conductivity of Doped Ceria Electrolytes

## Abstract

This paper studies the ionic conductivity of Gd-doped ceria (GDC) electrolytes for
solid oxide fuel cells. The conductivity reaches 0.12 S cm^-1 at 600 °C. We find
that the activation energy E_a is 0.82 eV.

**Keywords:** solid oxide fuel cell; Gd-doped ceria; ionic conductivity

## 1. Introduction

Doped ceria is a promising electrolyte intermediate-temperature fuel cells [1].
Its lattice fluorite structure is stable below 700 °C, but Ce^4+ reduction limits
operation.

### 1.1 Motivation

The oxygen vacancy concentration δ is tuned by dopant content.

## 2. Experimental

Powders were prepared by glycine nitrate process. Table 1 summarizes samples.

Table 1. Composition and conductivity of samples.

| Sample | Dopant (mol%) | σ at 600 °C (S cm^-1) | E_a (eV) |
|:------|:-------------:|---------------------:|--------:|
| GDC10 | 10 | 0.085 | 0.88 |
| GDC20 | 20 | 0.120 | 0.82 |

## 3. Results

Conductivity increases with dopant up to 20 mol%, then decreases due to vacancy
ordering. The Raman peak at ~465 cm^-1 confirms the fluorite phase.
"""

REFS = [
    {
        "title": "Electrical properties of Gd-doped ceria",
        "authors": "M. Mogensen, N. M. Sammes",
        "year": 2000,
        "source": "Solid State Ionics",
        "doi": "10.1016/S0167-2738(00)00497-8",
    },
    {
        "title": "Oxide ion conducting electrolytes for SOFCs",
        "authors": "B. C. H. Steele",
        "year": 1999,
        "source": "Journal of Materials Science",
        "doi": "",
    },
]


def _make_paper(cite_style: str = "gb7714") -> dict:
    return {
        "title": "Ionic Conductivity of Doped Ceria Electrolytes",
        "stages": {
            "write": {"content": WRITE_CONTENT},
            "search": {"refs": REFS},
        },
    }


def test_export_returns_docx_bytes():
    data = export_docx(_make_paper(), cite_style="gb7714")
    assert data is not None
    # docx 是 zip，魔数 PK
    assert data[:2] == b"PK"
    # 能被 python-docx 重新打开
    doc = Document(io.BytesIO(data))
    assert len(doc.paragraphs) > 0


def test_docx_structure():
    data = export_docx(_make_paper(), cite_style="gb7714")
    assert data is not None
    doc = Document(io.BytesIO(data))

    paras = doc.paragraphs
    styles = [p.style.name for p in paras]
    texts = [p.text for p in paras]
    full = "\n".join(texts)

    # 存在 Heading 段落（Heading 1 / Heading 2）
    heading_styles = [s for s in styles if s.startswith("Heading")]
    assert heading_styles, "应有 Word 内置 Heading 段落"
    assert "Heading 1" in styles
    assert "Heading 2" in styles

    # 论文标题（居中加粗，17pt）——非 Heading，但文本等于合成标题
    assert "Ionic Conductivity of Doped Ceria Electrolytes" in full

    # Abstract 标识
    assert "Abstract" in full
    # Keywords 引导
    assert "Keywords" in full
    # 关键词内容
    assert "solid oxide fuel cell" in full

    # 章节
    assert "Introduction" in full
    assert "Experimental" in full
    assert "Results" in full

    # 表格：至少 1 个表格，且有 4 列 3 行（表头 + 2 数据行）
    assert len(doc.tables) >= 1
    tbl = doc.tables[0]
    assert len(tbl.rows) == 3
    assert len(tbl.columns) == 4
    header_text = " ".join(c.text for c in tbl.rows[0].cells)
    assert "Dopant" in header_text

    # 参考文献条目：[1] / [2] 出现
    assert "[1]" in full
    assert "[2]" in full
    assert "Mogensen" in full

    # 全文无任何工作流标签文字
    for bad in ("选题设定", "文献检索", "大纲生成", "逐章写作", "质量巡检",
                "模拟审稿", "引用说明"):
        assert bad not in full, f"工作流标签残留: {bad}"


def test_no_workflow_leak_in_parsed_blocks():
    """解析后的块里不应含过程块（引用说明等）。"""
    parsed = parse_paper(WRITE_CONTENT, "fallback")
    full_blocks = "\n".join(getattr(b, "raw", "") for b in parsed.blocks)
    for bad in ("引用说明", "选题设定", "文献检索"):
        assert bad not in full_blocks


def test_empty_write_returns_none():
    p = {"title": "x", "stages": {"write": {"content": "  "}, "search": {"refs": []}}}
    assert export_docx(p) is None


def test_cite_style_difference_gb7714_vs_acs():
    """gb7714 与 acs 排版应产生不同的参考文献行。"""
    p = _make_paper()
    gb = export_docx(p, cite_style="gb7714")
    acs = export_docx(p, cite_style="acs")
    assert gb is not None and acs is not None
    # 直接对比 ref_lines 的文本差异（docx 内容同源）
    gb_lines = ref_lines(REFS, "gb7714")
    acs_lines = ref_lines(REFS, "acs")
    assert gb_lines != acs_lines
    # gb7714 含 [J] 文献类型标识
    assert any("[J]" in ln for ln in gb_lines)
    # acs 不含 [J]
    assert not any("[J]" in ln for ln in acs_lines)


def test_docx_zip_integrity():
    data = export_docx(_make_paper())
    assert data is not None
    zf = zipfile.ZipFile(io.BytesIO(data))
    names = zf.namelist()
    assert "word/document.xml" in names
    bad = zf.testzip()
    assert bad is None


# ---------------------------------------------------------------------------
# 插图（figures_dir）用例
# ---------------------------------------------------------------------------
WRITE_WITH_FIG_CAPTION = """# Doped Ceria Electrolytes

## Abstract

Test abstract.

**Keywords:** a; b

## 1. Introduction

Some text.

Figure 1. Schematic cross-section of the anode-supported cell.

More text.
"""


def test_export_without_figures_dir_no_images():
    """figures_dir=None 时 docx 不含内嵌图片媒体。"""
    p = {"title": "T", "stages": {"write": {"content": WRITE_WITH_FIG_CAPTION},
                                   "search": {"refs": []}}}
    data = export_docx(p)
    assert data is not None
    zf = zipfile.ZipFile(io.BytesIO(data))
    media = [n for n in zf.namelist() if n.startswith("word/media/")]
    assert media == []


def test_export_with_figures_dir_inserts_image(tmp_path):
    """figures_dir 含对应 PNG 时，docx 内嵌图片媒体。"""
    from matlens_kit.paper.fig_gen import gen_figures
    figs_dir = str(tmp_path / "figs")
    gen_figures(WRITE_WITH_FIG_CAPTION, figs_dir)

    p = {"title": "T", "stages": {"write": {"content": WRITE_WITH_FIG_CAPTION},
                                   "search": {"refs": []}}}
    data = export_docx(p, figures_dir=figs_dir)
    assert data is not None
    assert data[:2] == b"PK"

    zf = zipfile.ZipFile(io.BytesIO(data))
    media = [n for n in zf.namelist() if n.startswith("word/media/")]
    assert len(media) >= 1, f"expected embedded image, got {media}"

    # 图注文字仍在
    doc = Document(io.BytesIO(data))
    full = "\n".join(p.text for p in doc.paragraphs)
    assert "Figure 1." in full

